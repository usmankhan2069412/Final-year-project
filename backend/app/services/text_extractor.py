import asyncio
import csv
import re
import logging
from dataclasses import dataclass
from typing import Dict, Any, List
import httpx

import docx
import fitz  # PyMuPDF

from app.services.firecrawl_service import FirecrawlService

logger = logging.getLogger(__name__)


@dataclass
class CrawlResult:
    text: str
    pages_crawled: int
    total_chars: int
    crawl_duration_secs: float


class TextExtractor:
    """
    Unified text extraction service for documents (PDF, DOCX, TXT, CSV)
    and web pages (via Firecrawl).
    """

    @staticmethod
    def extract_pdf(file_path: str) -> str:
        text = []
        with fitz.open(file_path) as doc:
            for page in doc:
                text.append(page.get_text())
        return "\n".join(text)

    @staticmethod
    def extract_docx(file_path: str) -> str:
        doc = docx.Document(file_path)
        lines = []
        for child in doc.element.body:
            if child.tag.endswith("p"):
                p = docx.text.paragraph.Paragraph(child, doc)
                if p.text.strip():
                    lines.append(p.text.strip())
            elif child.tag.endswith("tbl"):
                table = docx.table.Table(child, doc)
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        # Fallback if body iteration produced no lines
        if not lines:
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            table_lines = []
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        table_lines.append(" | ".join(cells))
            lines = paragraphs + table_lines
        return "\n".join(lines)

    @staticmethod
    def extract_txt(file_path: str) -> str:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    @staticmethod
    def extract_csv(file_path: str) -> str:
        text_rows = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            header = None
            for row in reader:
                cleaned = [col.strip() for col in row]
                if not any(cleaned):
                    continue
                if header is None:
                    header = cleaned
                    text_rows.append(" | ".join(header))
                else:
                    row_repr = "; ".join(f"{h}: {val}" for h, val in zip(header, cleaned) if val)
                    text_rows.append(row_repr if row_repr else " | ".join(cleaned))
        return "\n".join(text_rows)

    @staticmethod
    async def fast_validate_url(url: str) -> Dict[str, Any]:
        """
        Fast, lightweight reachability and title check (5s timeout)
        without consuming Firecrawl credits or causing UI lag.
        """
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        }
        try:
            async with httpx.AsyncClient(timeout=5.0, follow_redirects=True) as client:
                resp = await client.get(url, headers=headers)
                if resp.status_code >= 400 and resp.status_code not in (401, 403):
                    raise ValueError(f"Website returned HTTP status {resp.status_code}")
                
                title = ""
                html_chunk = resp.text[:32768] if resp.text else ""
                match = re.search(r"<title[^>]*>(.*?)</title>", html_chunk, re.IGNORECASE | re.DOTALL)
                if match:
                    title = match.group(1).strip()
                
                return {
                    "status": "ok",
                    "title": title or "Website Accessible",
                    "links_found": 0,
                    "chars_extracted": len(resp.text) if resp.text else 0,
                }
        except (httpx.ConnectError, httpx.TimeoutException) as e:
            raise ValueError(f"Could not connect to website within 5s: {e}")
        except Exception as e:
            if isinstance(e, ValueError):
                raise e
            raise ValueError(f"Website validation failed: {e}")

    @staticmethod
    def validate_url_sync(url: str) -> Dict[str, Any]:
        """
        Validate URL accessibility using Firecrawl.
        """
        return FirecrawlService.validate_url(url)

    @staticmethod
    def extract_url_sync(url: str) -> CrawlResult:
        """
        Extract clean Markdown and metadata from a website using Firecrawl.
        """
        res = FirecrawlService.crawl_site(url)
        return CrawlResult(
            text=res.text,
            pages_crawled=res.pages_crawled,
            total_chars=res.total_chars,
            crawl_duration_secs=res.crawl_duration_secs,
        )

    @staticmethod
    async def extract_url(url: str) -> CrawlResult:
        """
        Async wrapper using FirecrawlService.crawl_site_async.
        """
        res = await FirecrawlService.crawl_site_async(url)
        return CrawlResult(
            text=res.text,
            pages_crawled=res.pages_crawled,
            total_chars=res.total_chars,
            crawl_duration_secs=res.crawl_duration_secs,
        )
